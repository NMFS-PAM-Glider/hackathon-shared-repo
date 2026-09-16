"""
Generic PowerPoint deck-building helpers - title/section/bullet/image slides on a 16:9 deck.
Not specific to this project; same pattern as the glider server-comparison plots notebook.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image

BULLET_FONT_SIZE = Pt(20)
TITLE_FONT_SIZE = Pt(24)
DESC_FONT_SIZE = Pt(12)


def new_presentation():
    """A blank 16:9 (13.333in x 7.5in) presentation."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def add_title_slide(prs, title, subtitle=''):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle:
        slide.placeholders[1].text = subtitle
    return slide


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[2])
    slide.shapes.title.text = title
    return slide


def add_bullet_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    body.word_wrap = True
    body.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    body.text = bullets[0]
    body.paragraphs[0].font.size = BULLET_FONT_SIZE
    for b in bullets[1:]:
        p = body.add_paragraph()
        p.text = b
        p.font.size = BULLET_FONT_SIZE
    return slide


def add_picture_fit(slide, img_path, left, top, max_width, max_height):
    """Add an image to `slide`, scaled to fit inside the given box while preserving aspect ratio."""
    img_w_px, img_h_px = Image.open(img_path).size
    aspect = img_w_px / img_h_px
    box_aspect = max_width / max_height
    if aspect > box_aspect:
        width, height = max_width, int(max_width / aspect)
    else:
        height, width = max_height, int(max_height * aspect)
    left_offset = left + (max_width - width) // 2
    top_offset = top + (max_height - height) // 2
    slide.shapes.add_picture(img_path, left_offset, top_offset, width=width, height=height)


def add_image_slide(prs, title, img_path, description):
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.6))
    tb.text_frame.text = title
    tb.text_frame.paragraphs[0].font.size = TITLE_FONT_SIZE
    tb.text_frame.paragraphs[0].font.bold = True
    add_picture_fit(slide, img_path, Inches(0.5), Inches(1.0), Inches(12.3), Inches(5.0))
    desc = slide.shapes.add_textbox(Inches(0.5), Inches(6.2), Inches(12.3), Inches(1.1))
    desc.text_frame.word_wrap = True
    desc.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    desc.text_frame.text = description
    desc.text_frame.paragraphs[0].font.size = DESC_FONT_SIZE
    return slide
